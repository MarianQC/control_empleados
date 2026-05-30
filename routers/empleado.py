# Endpoints del empleado: marcar entrada/salida con dos turnos y sucursal

from datetime import date, datetime
from fastapi import APIRouter, Cookie, HTTPException, Request, Query, Form
from fastapi.responses import RedirectResponse, StreamingResponse
from fastapi.templating import Jinja2Templates
from modelos import Asistencia, Empleado, SessionLocal, Log

import io

templates = Jinja2Templates(directory="templates")
router = APIRouter()


@router.get("/dashboard/empleado")
def dashboard_empleado(request: Request, cedula: str = Cookie(None)):
    if not cedula:
        return RedirectResponse(url="/login", status_code=303)

    db = SessionLocal()
    from modelos import Sucursal
    
    empleado = db.query(Empleado).filter(Empleado.cedula == cedula).first()
    asistencias = empleado.asistencias

    # ── CAMBIO: forzar carga de la relación sucursal antes de cerrar la sesión
    for a in asistencias:
        _ = a.sucursal  # esto fuerza el lazy load mientras la sesión está abierta

    sucursales_activas = db.query(Sucursal).filter(Sucursal.activo == True).all()
    db.close()

    return templates.TemplateResponse(
        request=request,
        name="empleado/dashboard.html",
        context={
            "request": request,
            "empleado": empleado,
            "asistencias": asistencias,
            "sucursales": sucursales_activas,
            "hoy": date.today()
        }
    )


# ── Turno 1: Entrada ──────────────────────────────────────────
@router.post("/asistencia/entrada")
def marcar_entrada(
    request: Request,
    cedula: str = Cookie(None),
    cedula_confirmacion: str = Form(...),
    sucursal_id: int = Form(...)
):
    if not cedula:
        return RedirectResponse(url="/login", status_code=303)

    if cedula != cedula_confirmacion:
        raise HTTPException(status_code=400, detail="La cédula de confirmación no coincide")

    db = SessionLocal()
    hoy = date.today()

    reg = db.query(Asistencia).filter(
        Asistencia.cedula_empleado == cedula,
        Asistencia.fecha_entrada_1 == hoy
    ).first()

    if reg and reg.hora_entrada_1:
        db.close()
        raise HTTPException(status_code=400, detail="Ya marcó entrada del turno 1 hoy")

    ahora = datetime.now().replace(microsecond=0).time()

    if reg:
        reg.hora_entrada_1 = ahora
        reg.sucursal_id    = sucursal_id
    else:
        db.add(Asistencia(
            fecha_entrada_1 = hoy,
            hora_entrada_1  = ahora,
            cedula_empleado = cedula,
            sucursal_id     = sucursal_id
        ))

    db.commit()
    db.close()
    return RedirectResponse(url="/dashboard/empleado", status_code=303)


# ── Turno 1: Salida ───────────────────────────────────────────
@router.post("/asistencia/salida")
def marcar_salida(
    request: Request,
    cedula: str = Cookie(None),
    cedula_confirmacion: str = Form(...),
):
    if not cedula:
        return RedirectResponse(url="/login", status_code=303)

    if cedula != cedula_confirmacion:
        raise HTTPException(status_code=400, detail="La cédula de confirmación no coincide")

    db = SessionLocal()
    hoy = date.today()

    reg = db.query(Asistencia).filter(
        Asistencia.cedula_empleado == cedula,
        Asistencia.fecha_entrada_1 == hoy
    ).first()

    if not reg or not reg.hora_entrada_1:
        db.close()
        raise HTTPException(status_code=400, detail="No ha marcado entrada del turno 1 hoy")

    if reg.hora_salida_1:
        db.close()
        raise HTTPException(status_code=400, detail="Ya marcó salida del turno 1 hoy")

    reg.fecha_salida_1 = hoy
    reg.hora_salida_1  = datetime.now().replace(microsecond=0).time()
    db.commit()
    db.close()
    return RedirectResponse(url="/dashboard/empleado", status_code=303)


# ── Turno 2: Entrada ──────────────────────────────────────────
@router.post("/asistencia/entrada2")
def marcar_entrada2(
    request: Request,
    cedula: str = Cookie(None),
    cedula_confirmacion: str = Form(...),
    sucursal_id: int = Form(...)
):
    if not cedula:
        return RedirectResponse(url="/login", status_code=303)

    if cedula != cedula_confirmacion:
        raise HTTPException(status_code=400, detail="La cédula de confirmación no coincide")

    db = SessionLocal()
    hoy = date.today()

    reg = db.query(Asistencia).filter(
    Asistencia.cedula_empleado == cedula,
    Asistencia.hora_salida_1 != None,
    Asistencia.hora_entrada_2 == None).order_by(Asistencia.id.desc()).first()

    if not reg:
        db.close()
        raise HTTPException(status_code=400, detail="Debe marcar el turno 1 antes del turno 2")

    if not reg.hora_salida_1:
        db.close()
        raise HTTPException(status_code=400, detail="Debe marcar salida del turno 1 antes de iniciar el turno 2")

    if reg.hora_entrada_2:
        db.close()
        raise HTTPException(status_code=400, detail="Ya marcó entrada del turno 2 hoy")

    reg.fecha_entrada_2 = hoy
    reg.hora_entrada_2  = datetime.now().replace(microsecond=0).time()
    reg.sucursal_id     = sucursal_id
    db.commit()
    db.close()
    return RedirectResponse(url="/dashboard/empleado", status_code=303)


# ── Turno 2: Salida ───────────────────────────────────────────
@router.post("/asistencia/salida2")
def marcar_salida2(
    request: Request,
    cedula: str = Cookie(None),
    cedula_confirmacion: str = Form(...),
):
    if not cedula:
        return RedirectResponse(url="/login", status_code=303)

    if cedula != cedula_confirmacion:
        raise HTTPException(status_code=400, detail="La cédula de confirmación no coincide")

    db = SessionLocal()
    hoy = date.today()

    reg = db.query(Asistencia).filter(
    Asistencia.cedula_empleado == cedula,
    Asistencia.hora_entrada_2 != None,
    Asistencia.hora_salida_2 == None).order_by(Asistencia.id.desc()).first()

    if not reg or not reg.hora_entrada_2:
        db.close()
        raise HTTPException(status_code=400, detail="No ha marcado entrada del turno 2 hoy")

    if reg.hora_salida_2:
        db.close()
        raise HTTPException(status_code=400, detail="Ya marcó salida del turno 2 hoy")

    reg.fecha_salida_2 = hoy
    reg.hora_salida_2  = datetime.now().replace(microsecond=0).time()
    db.commit()
    db.close()
    return RedirectResponse(url="/dashboard/empleado", status_code=303)


# ── Exportar Excel ────────────────────────────────────────────
@router.get("/exportar/asistencia/excel")
def exportar_asistencia_excel(
    cedula: str = Cookie(None),
    cedula_query: str = Query(None, alias="cedula")
):
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill

    cedula_final = cedula or cedula_query
    if not cedula_final:
        raise HTTPException(status_code=401, detail="No autorizado")

    db = SessionLocal()
    registros = (
        db.query(Asistencia)
        .filter(Asistencia.cedula_empleado == cedula_final)
        .order_by(Asistencia.fecha_entrada_1.desc())
        .all()
    )
    empleado = db.query(Empleado).filter(Empleado.cedula == cedula_final).first()
    db.close()

    if not empleado:
        raise HTTPException(status_code=404, detail="Empleado no encontrado")

    wb  = openpyxl.Workbook()
    ws  = wb.active
    ws.title = "Asistencia"

    fuente_titulo  = Font(name="Calibri", bold=True, size=14, color="FFFFFF")
    fuente_header  = Font(name="Calibri", bold=True, size=11, color="FFFFFF")
    fuente_info    = Font(name="Calibri", size=10)
    fuente_bold    = Font(name="Calibri", bold=True, size=10)
    relleno_titulo = PatternFill("solid", fgColor="1F3864")
    relleno_header = PatternFill("solid", fgColor="2E75B6")
    relleno_info   = PatternFill("solid", fgColor="D9E1F2")
    centro         = Alignment(horizontal="center", vertical="center")

    ws.merge_cells("A1:I1")
    ws["A1"] = "REPORTE DE ASISTENCIA"
    ws["A1"].font = fuente_titulo; ws["A1"].fill = relleno_titulo; ws["A1"].alignment = centro
    ws.row_dimensions[1].height = 28

    info_rows = [
        ("Empleado:",     f"{empleado.nombre} {empleado.apellido}"),
        ("Cédula:",       empleado.cedula),
        ("Departamento:", empleado.departamento),
        ("Generado el:",  datetime.now().strftime("%d/%m/%Y %H:%M")),
    ]
    for i, (label, valor) in enumerate(info_rows, start=2):
        ws.cell(row=i, column=1, value=label).font = fuente_bold
        ws.cell(row=i, column=1).fill             = relleno_info
        ws.cell(row=i, column=2, value=valor).font = fuente_info
        ws.cell(row=i, column=2).fill             = relleno_info

    headers = ["ID", "F.Entrada T1", "H.Entrada T1", "F.Salida T1", "H.Salida T1",
               "F.Entrada T2", "H.Entrada T2", "F.Salida T2", "H.Salida T2"]
    for col, h in enumerate(headers, start=1):
        cell = ws.cell(row=7, column=col, value=h)
        cell.font = fuente_header; cell.fill = relleno_header; cell.alignment = centro

    def fmt_date(d): return d.strftime("%d/%m/%Y") if d else "—"
    def fmt_time(t): return t.strftime("%H:%M:%S") if t else "—"

    for fila, r in enumerate(registros, start=8):
        datos = [
            r.id,
            fmt_date(r.fecha_entrada_1), fmt_time(r.hora_entrada_1),
            fmt_date(r.fecha_salida_1),  fmt_time(r.hora_salida_1),
            fmt_date(r.fecha_entrada_2), fmt_time(r.hora_entrada_2),
            fmt_date(r.fecha_salida_2),  fmt_time(r.hora_salida_2),
        ]
        for col, valor in enumerate(datos, start=1):
            cell = ws.cell(row=fila, column=col, value=valor)
            cell.font = fuente_info; cell.alignment = centro

    for col, ancho in enumerate([6,14,14,14,14,14,14,14,14], start=1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = ancho

    output = io.BytesIO()
    wb.save(output); output.seek(0)

    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename=asistencia_{empleado.cedula}_{date.today()}.xlsx"}
    )